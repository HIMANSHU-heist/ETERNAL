"""
Converts non-tabular uploads (PDF, DOCX, TXT) into a CSV so the rest
of the pipeline (file_loader, registry, agents) never needs to know
the original format existed.
"""
from pathlib import Path
import pandas as pd

CONVERTIBLE_EXTENSIONS = {".pdf", ".docx", ".txt"}


class ConversionError(Exception):
    pass


def convert_to_csv(file_path: str, output_dir: Path) -> Path:
    path = Path(file_path)
    ext = path.suffix.lower()

    try:
        if ext == ".pdf":
            df = _pdf_to_df(path)
        elif ext == ".docx":
            df = _docx_to_df(path)
        elif ext == ".txt":
            df = _txt_to_df(path)
        else:
            raise ConversionError(f"'{ext}' is not convertible.")
    except ConversionError:
        raise
    except Exception as e:
        raise ConversionError(f"Failed to parse '{ext}' file: {e}") from e

    if df.empty:
        raise ConversionError("No tabular data could be extracted from this file.")

    out_path = output_dir / f"{path.stem}_converted.csv"
    df.to_csv(out_path, index=False)
    return out_path


def _pdf_to_df(path: Path) -> pd.DataFrame:
    import pdfplumber

    dfs = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables():
                if not table or len(table) < 2:
                    continue
                header, *data = table
                dfs.append(pd.DataFrame(data, columns=header))

    if not dfs:
        return pd.DataFrame()

    reference_columns = list(dfs[0].columns)
    for i, df in enumerate(dfs[1:], start=1):
        if list(df.columns) != reference_columns:
            raise ConversionError(
                "PDF contains tables with different columns; cannot safely merge them."
            )

    return pd.concat(dfs, ignore_index=True)


def _docx_to_df(path: Path) -> pd.DataFrame:
    import docx

    doc = docx.Document(path)
    if not doc.tables:
        return pd.DataFrame()

    dfs = []
    for table in doc.tables:
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        if len(rows) < 2:
            continue
        header, *data = rows
        dfs.append(pd.DataFrame(data, columns=header))

    if not dfs:
        return pd.DataFrame()

    reference_columns = list(dfs[0].columns)
    for i, df in enumerate(dfs[1:], start=1):
        if list(df.columns) != reference_columns:
            raise ConversionError(
                "DOCX contains tables with different columns; cannot safely merge them."
            )

    return pd.concat(dfs, ignore_index=True)


def _txt_to_df(path: Path) -> pd.DataFrame:
    # assumes delimited text (csv-like content with wrong extension)
    return pd.read_csv(path, sep=None, engine="python")